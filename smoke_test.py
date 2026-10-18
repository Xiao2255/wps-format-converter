"""Integration checks using generated fixtures only; never opens user documents."""
from pathlib import Path
from converters import convert
from PIL import Image
import pymupdf
from docx import Document


def main():
    root = Path('tmp/smoke').resolve()
    root.mkdir(parents=True, exist_ok=True)
    source = root / '透明图片.png'
    Image.new('RGBA', (40, 30), (0, 0, 0, 0)).save(source)
    for fmt in ['JPEG', 'PNG', 'WEBP', 'BMP', 'TIFF']:
        output = convert(source, root, '图片互转', fmt)
        with Image.open(output) as img:
            assert img.size == (40, 30)
            if fmt in ['JPEG', 'BMP']:
                assert img.getpixel((0, 0)) == (255, 255, 255)
    assert Image.open(source).mode == 'RGBA'
    print('PASS images, transparency, same-name protection', flush=True)
    doc = Document()
    doc.add_heading('文件转换测试', 0)
    doc.add_paragraph('这是一段可编辑的中文文字。Conversion test 123.')
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    for cell, value in zip([c for r in table.rows for c in r.cells], ['项目', '数量', '图片', '3']):
        cell.text = value
    word = root / '中文样例.docx'
    doc.save(word)
    pdf = convert(word, root, 'Word → PDF')
    with pymupdf.open(pdf) as result:
        assert 'Conversion test 123' in ''.join(p.get_text() for p in result)
        result[0].get_pixmap().save(str(root / 'wps-preview.png'))
    print('PASS WPS export', flush=True)
    back = convert(pdf, root, 'PDF → Word')
    edited = Document(back)
    text = ''.join(p.text for p in edited.paragraphs) + ''.join(c.text for t in edited.tables for r in t.rows for c in r.cells)
    assert 'Conversion test 123' in text
    assert all(value in text for value in ['可编辑', '项目', '数量', '图片'])
    assert edited.tables
    print('PASS PDF to editable DOCX with Chinese and table', flush=True)
    blank = root / 'blank.pdf'
    with pymupdf.open() as empty:
        empty.new_page()
        empty.save(blank)
    try:
        convert(blank, root, 'PDF → Word')
    except ValueError:
        print('PASS no-text PDF rejected', flush=True)
    else:
        raise AssertionError('Expected rejection')


if __name__ == '__main__':
    main()
